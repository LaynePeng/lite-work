# SPDX-License-Identifier: Apache-2.0
# Copyright (c) 2026 lite-work contributors

"""Office 工具单测：docx_append 迭代写作 + data_analyze 文件直读。"""
from __future__ import annotations

import asyncio
import json
import os

import pytest

from litework.tools.office import OfficeTools


def _run(tool: str, args: dict) -> str:
    tools = OfficeTools(os.getcwd())
    return asyncio.run(tools.execute(tool, args))


# ---------------------------------------------------------------- docx_append

def test_docx_append_roundtrip(tmp_path, monkeypatch):
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    # 1) 先创建
    r1 = asyncio.run(call("docx_create", {
        "content": "# 方案\n\n第一版内容",
        "filename": "方案.docx",
    }))
    assert "方案.docx" in r1

    # 2) 追加章节
    r2 = asyncio.run(call("docx_append", {
        "path": "产出物/方案.docx",
        "content": "\n# 补充章节\n\n追加的正文段落",
    }))
    assert "已追加内容" in r2

    # 3) 校验：原内容与新内容都在
    from docx import Document
    doc = Document(str(tmp_path / "产出物" / "方案.docx"))
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)
    assert "第一版内容" in joined
    assert "追加的正文段落" in joined
    # 标题段落存在（补充章节渲染为 heading）
    assert any("补充章节" in t for t in texts)


def test_docx_append_page_break_and_errors(tmp_path):
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    asyncio.run(call("docx_create", {"content": "初始", "filename": "d.docx"}))

    # 分页符参数可用
    r = asyncio.run(call("docx_append", {
        "path": "产出物/d.docx", "content": "新页内容", "page_break": True,
    }))
    assert "已追加内容" in r

    # 文件不存在
    r2 = asyncio.run(call("docx_append", {"path": "产出物/nope.docx", "content": "x"}))
    assert "不存在" in r2
    # 非 docx 扩展名
    r3 = asyncio.run(call("docx_append", {"path": "产出物/x.txt", "content": "x"}))
    assert "仅支持" in r3
    # 空内容
    r4 = asyncio.run(call("docx_append", {"path": "产出物/d.docx", "content": "  "}))
    assert "为空" in r4
    # 路径越界
    r5 = asyncio.run(call("docx_append", {"path": "../../etc/hosts", "content": "x"}))
    assert "越界" in r5 or "不存在" in r5 or "仅支持" in r5


# ---------------------------------------------------------------- 工具执行不阻塞事件循环（回归守护）

def test_office_execute_does_not_block_event_loop(tmp_path):
    """office 工具为同步重活（matplotlib/subprocess），必须放线程池——
    直接同步调用会阻塞 asyncio 事件循环，曾导致 SSE/审批全部冻结
    （用户侧表现为应用完全卡死）。此测试为该回归的守护。
    """
    import asyncio
    import time

    tools = OfficeTools(str(tmp_path))
    ticks = []

    async def heartbeat():
        # 事件循环心跳：每 50ms 记录一次，工具执行期间必须持续跳动
        for _ in range(20):
            ticks.append(time.monotonic())
            await asyncio.sleep(0.05)

    async def run_chart():
        return await tools.execute("chart_make", {
            "data": '{"labels": ["A", "B"], "values": [1, 2]}',
            "chart_type": "bar",
            "filename": "hb.png",
        })

    async def main():
        hb = asyncio.create_task(heartbeat())
        await run_chart()
        await hb

    asyncio.run(main())

    # 心跳间隔应基本均匀（最大间隔 < 1s 即视为未阻塞；matplotlib 出图
    # 本身耗时无妨——它在别的线程，事件循环必须保持响应）
    gaps = [b - a for a, b in zip(ticks, ticks[1:])]
    assert gaps, "心跳未运行"
    assert max(gaps) < 1.0, f"事件循环被阻塞：心跳间隔 {max(gaps):.2f}s"
    assert (tmp_path / "产出物" / "hb.png").exists()


# ---------------------------------------------------------------- 图表代码块自动渲染（工具层兜底）

def test_docx_plantuml_block_auto_render_or_keep(tmp_path):
    """content 中的 ```plantuml 代码块：有引擎→渲染成图嵌入；无引擎→保留源码文本（不丢）。"""
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    content = "# 架构说明\n\n```plantuml\n@startuml\nAlice -> Bob: hi\n@enduml\n```\n\n正文结束"
    r = asyncio.run(call("docx_create", {"content": content, "filename": "带图.docx"}))
    assert "带图.docx" in r

    from docx import Document
    doc = Document(str(tmp_path / "产出物" / "带图.docx"))
    # 无论渲染成功与否，源码文本或图片至少保留其一：
    # - 渲染成功：inline_shapes >= 1
    # - 无引擎回退：段落里保留 @startuml 源码
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert (len(doc.inline_shapes) >= 1) or ("@startuml" in all_text), \
        "plantuml 代码块既没渲染成图也没保留源码——内容丢失！"

    # diagrams 缓存目录只在渲染成功时存在
    if (tmp_path / "产出物" / "diagrams").is_dir():
        pngs = list((tmp_path / "产出物" / "diagrams").glob("*.png"))
        assert pngs, "diagrams 目录存在但没有 PNG"


def test_docx_mermaid_block_kept_when_render_fails(tmp_path, monkeypatch):
    """渲染失败（引擎缺失/语法错）时必须回退保留源码文本——内容不丢是硬约束。"""
    tools = OfficeTools(str(tmp_path))

    # 强制渲染失败：monkeypatch _render_diagram_block 返回 None
    monkeypatch.setattr(
        OfficeTools, "_render_diagram_block",
        lambda self, lang, code, seq: None,
    )

    async def call(name, args):
        return await tools.execute(name, args)

    content = "```mermaid\nflowchart LR\nA-->B\n```\n"
    r = asyncio.run(call("docx_create", {"content": content, "filename": "兜底.docx"}))
    assert "兜底.docx" in r

    from docx import Document
    doc = Document(str(tmp_path / "产出物" / "兜底.docx"))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "flowchart LR" in all_text, "渲染失败时源码文本被丢弃了！"
    assert "A-->B" in all_text


def test_docx_normal_code_block_unaffected(tmp_path):
    """普通代码块（python 等）行为不变：仍作为等宽文本渲染。"""
    tools = OfficeTools(str(tmp_path))

    async def call():
        return await tools.execute("docx_create", {
            "content": "```python\nprint('hi')\n```",
            "filename": "代码.docx",
        })

    asyncio.run(call())
    from docx import Document
    doc = Document(str(tmp_path / "产出物" / "代码.docx"))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "print('hi')" in all_text
    assert len(doc.inline_shapes) == 0  # 普通代码块不应被渲染成图


# ---------------------------------------------------------------- data_analyze 文件直读

def test_data_analyze_xlsx_path(tmp_path):
    # 准备一个 xlsx 文件（模拟用户上传）
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["城市", "销售额"])
    ws.append(["北京", 100])
    ws.append(["上海", 200])
    uploads = tmp_path / "素材"
    uploads.mkdir()
    wb.save(str(uploads / "销售.xlsx"))

    tools = OfficeTools(str(tmp_path))

    async def call():
        return await tools.execute("data_analyze", {
            "path": "素材/销售.xlsx",
            "instructions": "描述性统计",
        })

    r = asyncio.run(call())
    assert "数据分析结果" in r
    assert "2 行" in r  # 2 行 × 2 列
    assert "城市" in r


def test_data_analyze_csv_path_and_sheet_error(tmp_path):
    csv_path = tmp_path / "素材" / "数据.csv"
    csv_path.parent.mkdir()
    csv_path.write_text("名称,数量\n甲,1\n乙,2\n", encoding="utf-8")

    tools = OfficeTools(str(tmp_path))

    async def call(args):
        return await tools.execute("data_analyze", args)

    # CSV 直读
    r = asyncio.run(call({"path": "素材/数据.csv", "instructions": "概览"}))
    assert "数据分析结果" in r
    assert "甲" in r or "名称" in r

    # 文件不存在
    r2 = asyncio.run(call({"path": "素材/nope.csv", "instructions": "概览"}))
    assert "不存在" in r2
    # 不支持的扩展名
    (tmp_path / "素材" / "f.txt").write_text("x", encoding="utf-8")
    r3 = asyncio.run(call({"path": "素材/f.txt", "instructions": "概览"}))
    assert "不支持" in r3
    # 路径越界
    r4 = asyncio.run(call({"path": "../../etc/passwd", "instructions": "概览"}))
    assert "越界" in r4

    # data 参数仍然兼容（不带 path）
    r5 = asyncio.run(call({
        "data": "名称,数量\n甲,1\n乙,2\n", "instructions": "概览",
    }))
    assert "数据分析结果" in r5


# ---------------------------------------------------------------- 读取已有办公文件

def test_docx_read_roundtrip(tmp_path):
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    asyncio.run(call("docx_create", {
        "content": "# 标题\n\n正文段落\n\n| 列A | 列B |\n| --- | --- |\n| 1 | 2 |",
        "filename": "读.docx",
    }))
    r = asyncio.run(call("docx_read", {"path": "产出物/读.docx"}))
    assert "已读取" in r
    assert "标题" in r
    assert "正文段落" in r
    assert "列A" in r

    # 文件不存在
    r2 = asyncio.run(call("docx_read", {"path": "产出物/nope.docx"}))
    assert "不存在" in r2
    # 路径越界
    r3 = asyncio.run(call("docx_read", {"path": "../../etc/hosts"}))
    assert "越界" in r3 or "不存在" in r3


def test_xlsx_read_roundtrip(tmp_path):
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    asyncio.run(call("xlsx_create", {
        "data": '{"sheet1": [{"城市": "北京", "销量": 100}, {"城市": "上海", "销量": 200}]}',
        "filename": "读.xlsx",
    }))
    r = asyncio.run(call("xlsx_read", {"path": "产出物/读.xlsx"}))
    assert "已读取" in r
    assert "北京" in r
    assert "上海" in r
    assert "销量" in r

    # sheet 参数
    r2 = asyncio.run(call("xlsx_read", {"path": "产出物/读.xlsx", "sheet": "sheet1"}))
    assert "北京" in r2
    # 不存在的 sheet
    r3 = asyncio.run(call("xlsx_read", {"path": "产出物/读.xlsx", "sheet": "nope"}))
    assert "不存在" in r3


def test_pptx_read_roundtrip(tmp_path):
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    asyncio.run(call("pptx_create", {
        "slides": '[{"title": "第一页", "bullets": ["要点1", "要点2"]}]',
        "filename": "读.pptx",
    }))
    r = asyncio.run(call("pptx_read", {"path": "产出物/读.pptx"}))
    assert "已读取" in r
    assert "第一页" in r
    assert "要点1" in r


def test_pdf_read_roundtrip(tmp_path):
    tools = OfficeTools(str(tmp_path))

    async def call(name, args):
        return await tools.execute(name, args)

    # reportlab 默认字体不支持中文（会被丢弃），用 ASCII 验证提取链路
    asyncio.run(call("pdf_create", {"content": "# PDF Title\n\nPDF body content", "filename": "read.pdf"}))
    r = asyncio.run(call("pdf_read", {"path": "产出物/read.pdf"}))
    assert "PDF 共" in r
    assert "PDF Title" in r or "PDF body" in r


# ---------------------------------------------------------------- 本地插件加载器

def test_plugin_loader_loads_cordis_plugin(tmp_path):
    """~/.lite-work/plugins/ 下的 Cordis ToolPlugin 子类应被自动发现并实例化。"""
    from litework.tools.plugin_loader import load_plugins

    plugins_dir = tmp_path / "plugins"
    plugins_dir.mkdir()
    (plugins_dir / "my_tools.py").write_text(
        "from litework.tools.plugin import ToolPlugin\n"
        "from litework.core.types import ToolDefinition\n"
        "class HelloPlugin(ToolPlugin):\n"
        "    name = 'hello-plugin'\n"
        "    def get_tools(self):\n"
        "        return [ToolDefinition(name='hello', description='say hello', parameters={'type':'object','properties':{}})]\n"
        "    async def execute(self, name, args):\n"
        "        return 'hello from plugin'\n",
        encoding="utf-8",
    )
    plugins = load_plugins(str(tmp_path))
    assert any(p.name == "hello-plugin" for p in plugins)


def test_plugin_loader_ignores_bad_modules(tmp_path):
    """加载失败的插件不应阻断其他插件。"""
    from litework.tools.plugin_loader import load_plugins

    plugins_dir = tmp_path / "plugins"
    plugins_dir.mkdir()
    (plugins_dir / "broken.py").write_text("raise ValueError('boom')\n", encoding="utf-8")
    (plugins_dir / "empty.py").write_text("x = 1\n", encoding="utf-8")
    plugins = load_plugins(str(tmp_path))
    assert plugins == []


def test_plugin_removed_tools_unregisters(tmp_path):
    """插件 removed_tools 声明删除的工具应从 registry 移除。"""
    from litework.core.kernel import Kernel
    from litework.tools.plugin import ToolPlugin
    from litework.tools.registry import ToolRegistry

    registry = ToolRegistry()
    registry.register("builtin_tool", "desc", {"type": "object", "properties": {}},
                      lambda args: "builtin")

    class RemovePlugin(ToolPlugin):
        name = "remove-plugin"
        removed_tools = ["builtin_tool"]

        def get_tools(self):
            from litework.core.types import ToolDefinition
            return [ToolDefinition(name="my_tool", description="mine",
                                   parameters={"type": "object", "properties": {}})]

        async def execute(self, name, args):
            return "ok"

    kernel = Kernel(session_id="test")
    kernel.register_service("tools", registry)
    kernel.use(RemovePlugin())
    assert not registry.has("builtin_tool"), "removed_tools 未生效"
    assert registry.has("my_tool")


def test_plugin_import_overwrite(tmp_path):
    """同名插件已存在时 overwrite=True 先删后装（更新）。"""
    from litework.tools.plugin_loader import import_zip_bytes
    import io
    import zipfile

    def make_zip(body: str) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("myplug/plugin.py", body)
        return buf.getvalue()

    v1 = make_zip(
        "from litework.tools.plugin import ToolPlugin\n"
        "class P(ToolPlugin):\n"
        "    name='myplug'\n"
        "    def get_tools(self):\n"
        "        from litework.core.types import ToolDefinition\n"
        "        return [ToolDefinition(name='t1', description='v1', parameters={'type':'object','properties':{}})]\n"
        "    async def execute(self, name, args): return 'v1'\n"
    )
    v2 = make_zip(
        "from litework.tools.plugin import ToolPlugin\n"
        "class P(ToolPlugin):\n"
        "    name='myplug'\n"
        "    def get_tools(self):\n"
        "        from litework.core.types import ToolDefinition\n"
        "        return [ToolDefinition(name='t2', description='v2', parameters={'type':'object','properties':{}})]\n"
        "    async def execute(self, name, args): return 'v2'\n"
    )

    r1 = import_zip_bytes(str(tmp_path), v1)
    assert r1[0]["name"] == "myplug"
    # 不覆盖 → 报错
    try:
        import_zip_bytes(str(tmp_path), v2)
        assert False, "应拒绝覆盖"
    except ValueError:
        pass
    # 覆盖 → 成功
    r2 = import_zip_bytes(str(tmp_path), v2, overwrite=True)
    assert r2[0]["name"] == "myplug"

    # 验证已替换为新版（tools 变了）
    from litework.tools.plugin_loader import list_plugins
    meta = list_plugins(str(tmp_path))
    assert any(p["name"] == "myplug" and "t2" in p["tools"] and "t1" not in p["tools"] for p in meta)


# ---------------------------------------------------------------- semver

def test_semver_parse():
    from litework.tools.plugin_loader import semver_parse, semver_compare
    assert semver_parse("1.2.3") == (1, 2, 3)
    assert semver_parse("v0.5.0") == (0, 5, 0)
    assert semver_parse("1.0.0-beta") == (1, 0, 0, "beta")
    assert semver_parse("") is None
    assert semver_parse("abc") is None
    assert semver_compare("1.0.0", "1.0.0") == 0
    assert semver_compare("1.0.0", "2.0.0") == -1
    assert semver_compare("2.0.0", "1.0.0") == 1
    # 预发布 < 正式版
    assert semver_compare("1.0.0-rc.1", "1.0.0") == -1


def test_installed_json_record(tmp_path):
    """installed.json 读写与删除后自动清理。"""
    from litework.tools.plugin_loader import (
        read_installed, record_installed, unrecord_installed,
    )
    cfg = str(tmp_path)
    record_installed(cfg, "my-tool", "1.0.0", "https://github.com/user/repo")
    data = read_installed(cfg)
    assert data["my-tool"]["version"] == "1.0.0"
    assert data["my-tool"]["source"] == "https://github.com/user/repo"
    unrecord_installed(cfg, "my-tool")
    assert "my-tool" not in read_installed(cfg)


# ---------------------------------------------------------------- 目录插件导入语义

PLUGIN_PY = (
    "from litework.tools.plugin import ToolPlugin\n"
    "from litework.core.types import ToolDefinition\n"
    "class P(ToolPlugin):\n"
    "    name='myplug'\n"
    "    def get_tools(self):\n"
    "        return [ToolDefinition(name='t1', description='v1', parameters={'type':'object','properties':{}})]\n"
    "    async def execute(self, name, args): return 'ok'\n"
)


def test_import_local_dir_that_is_plugin_dir(tmp_path):
    """导入目录本身是插件目录（直接含 plugin.py）→ 以目录名作为插件名安装。"""
    from litework.tools.plugin_loader import import_source, list_plugins

    src = tmp_path / "myplug"
    src.mkdir()
    (src / "plugin.py").write_text(PLUGIN_PY, encoding="utf-8")
    (src / "requirements.txt").write_text("# no deps\n", encoding="utf-8")

    cfg = tmp_path / "cfg"
    r = import_source(str(cfg), str(src))
    assert [x["name"] for x in r] == ["myplug"]
    # 目录插件：requirements.txt 等伴生文件保留
    installed = cfg / "plugins" / "myplug"
    assert (installed / "plugin.py").is_file()
    assert (installed / "requirements.txt").is_file()
    assert any(p["name"] == "myplug" for p in list_plugins(str(cfg)))


def test_import_zip_with_plugin_py_at_root(tmp_path):
    """zip 根直接是 plugin.py（文件夹内压缩）+ name → 归位为目录插件。"""
    import io
    import zipfile

    from litework.tools.plugin_loader import import_zip_bytes

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("plugin.py", PLUGIN_PY)
        zf.writestr("requirements.txt", "# no deps\n")

    cfg = tmp_path / "cfg"
    r = import_zip_bytes(str(cfg), buf.getvalue(), name="zipped-plug")
    assert [x["name"] for x in r] == ["zipped-plug"]
    installed = cfg / "plugins" / "zipped-plug"
    assert (installed / "plugin.py").is_file()
    assert (installed / "requirements.txt").is_file()


def test_plugin_wheels_extraction(tmp_path, monkeypatch):
    """插件自带 wheels/*.whl → 解压到 libs/ 并可被插件 import（打包态唯一依赖机制）。"""
    import io
    import sys as _sys
    import zipfile

    from litework.tools.plugin_loader import import_source, list_plugins

    # 造一个假 wheel：zip 根目录含 my_whl_pkg/__init__.py（whl 标准结构）
    whl_buf = io.BytesIO()
    with zipfile.ZipFile(whl_buf, "w") as zf:
        zf.writestr("my_whl_pkg/__init__.py", "VALUE = 'from-wheel'\n")
    whl_bytes = whl_buf.getvalue()

    src = tmp_path / "wheeled-plug"
    src.mkdir()
    wheels = src / "wheels"
    wheels.mkdir()
    (wheels / "my_whl_pkg-1.0.0-py3-none-any.whl").write_bytes(whl_bytes)
    # 插件代码 import wheel 里的包并暴露其值（验证 libs/ 已入 sys.path）
    (src / "plugin.py").write_text(
        "from litework.tools.plugin import ToolPlugin\n"
        "from litework.core.types import ToolDefinition\n"
        "from my_whl_pkg import VALUE\n"
        "class WheeledPlugin(ToolPlugin):\n"
        "    name='wheeled-plug'\n"
        "    def get_tools(self):\n"
        "        return [ToolDefinition(name='whl_tool', description=VALUE,\n"
        "                               parameters={'type':'object','properties':{}})]\n"
        "    async def execute(self, name, args): return VALUE\n",
        encoding="utf-8",
    )

    cfg = tmp_path / "cfg"
    import_source(str(cfg), str(src))

    installed = cfg / "plugins" / "wheeled-plug"
    assert (installed / "libs" / "my_whl_pkg" / "__init__.py").is_file(), "wheel 未解压到 libs/"
    assert (installed / "libs" / ".wheels-stamp").is_file(), "缺少解压 stamp"

    # 元信息读取会真实执行插件模块 → import my_whl_pkg 成功即证明 libs 生效
    meta = list_plugins(str(cfg))
    p = next(x for x in meta if x["name"] == "wheeled-plug")
    assert "whl_tool" in p["tools"]

    # 幂等：二次加载不重复解压（stamp 命中）
    from litework.tools.plugin_loader import _extract_wheels
    import time as _time
    before = (installed / "libs" / ".wheels-stamp").stat().st_mtime_ns
    _time.sleep(0.01)
    _extract_wheels(str(installed / "wheels"), str(installed / "libs"))
    assert (installed / "libs" / ".wheels-stamp").stat().st_mtime_ns == before, "stamp 未命中，重复解压"

    # 清理 sys.path（测试隔离）
    libs = str(installed / "libs")
    if libs in _sys.path:
        _sys.path.remove(libs)
    _sys.modules.pop("my_whl_pkg", None)

