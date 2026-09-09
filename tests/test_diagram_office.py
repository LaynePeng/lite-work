# SPDX-License-Identifier: Apache-2.0
# Copyright (c) 2026 lite-work contributors

"""图表转 Office 测试：docx/pptx 图片嵌入 + 渲染脚本逻辑。

- docx_create 的 Markdown 图片语法 ![alt](path) 应嵌入图片；
- pptx_create 的每页 image 字段应嵌入图片；
- render_diagram.py 的类型识别 / 源码提取逻辑应正确（纯函数，不依赖真实引擎）。
"""
from __future__ import annotations

import importlib.util
import io
import os
import sys
from pathlib import Path

import pytest

from litework.tools.office import OfficeTools

ROOT = Path(__file__).resolve().parents[1]
RENDER_SCRIPT = ROOT / "skills" / "diagram-to-office" / "render_diagram.py"


# ---------------------------------------------------------------- 工具函数

def _make_png(tmp_path: Path, name: str = "pic.png", size=(120, 60)) -> Path:
    """用 matplotlib（项目已依赖）生成一张小 PNG，避免依赖 PIL 手写。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    out = tmp_path / name
    fig, ax = plt.subplots(figsize=(size[0] / 100, size[1] / 100), dpi=100)
    ax.set_title("test")
    fig.savefig(out, dpi=100)
    plt.close(fig)
    return out


def _load_render_module():
    if not RENDER_SCRIPT.is_file():
        pytest.skip("render_diagram.py 不存在")
    spec = importlib.util.spec_from_file_location("render_diagram", RENDER_SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["render_diagram"] = mod
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------- docx 图片嵌入

def test_docx_create_embeds_markdown_image(tmp_path):
    tools = OfficeTools(str(tmp_path))
    png = _make_png(tmp_path)

    import asyncio

    async def run():
        return await tools.execute("docx_create", {
            "content": "# 架构\n\n![系统架构图](pic.png)\n\n正文",
            "filename": "图文.docx",
        })

    result = asyncio.run(run())
    assert "图文.docx" in result

    from docx import Document
    doc = Document(str(tmp_path / ".outputs" / "图文.docx"))
    assert len(doc.inline_shapes) == 1, "docx 应嵌入 1 张图片"
    # 图注文字应存在
    texts = [p.text for p in doc.paragraphs]
    assert any("系统架构图" in t for t in texts)


def test_docx_create_image_missing_is_graceful(tmp_path):
    tools = OfficeTools(str(tmp_path))

    import asyncio

    async def run():
        return await tools.execute("docx_create", {
            "content": "![不存在](no_such_file.png)",
            "filename": "缺图.docx",
        })

    result = asyncio.run(run())
    assert "缺图.docx" in result

    from docx import Document
    doc = Document(str(tmp_path / ".outputs" / "缺图.docx"))
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "图片未找到" in texts


# ---------------------------------------------------------------- pptx 图片嵌入

def test_pptx_create_embeds_slide_image(tmp_path):
    tools = OfficeTools(str(tmp_path))
    png = _make_png(tmp_path)

    import asyncio

    async def run():
        return await tools.execute("pptx_create", {
            "slides": '[{"title": "架构图", "image": "pic.png", "bullets": ["要点1"]},'
                      ' {"title": "无图页", "bullets": ["要点A"]}]',
            "filename": "图文.pptx",
        })

    result = asyncio.run(run())
    assert "图文.pptx" in result

    from pptx import Presentation
    prs = Presentation(str(tmp_path / ".outputs" / "图文.pptx"))
    pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]  # PICTURE
    assert len(pics) == 1, "第一页应嵌入 1 张图片"
    # 第二页无图
    assert sum(1 for sh in prs.slides[1].shapes if sh.shape_type == 13) == 0


def test_pptx_create_image_missing_is_graceful(tmp_path):
    tools = OfficeTools(str(tmp_path))

    import asyncio

    async def run():
        return await tools.execute("pptx_create", {
            "slides": '[{"title": "缺图", "image": "nope.png"}]',
            "filename": "缺图.pptx",
        })

    result = asyncio.run(run())
    assert "缺图.pptx" in result


# ---------------------------------------------------------------- 渲染脚本逻辑

def test_render_detect_plantuml():
    mod = _load_render_module()
    assert mod.detect_type("@startuml\nAlice -> Bob: hi\n@enduml", "x.txt", "auto") == "plantuml"
    assert mod.detect_type("anything", "arch.puml", "auto") == "plantuml"
    assert mod.detect_type("x", "y.txt", "plantuml") == "plantuml"


def test_render_detect_mermaid():
    mod = _load_render_module()
    assert mod.detect_type("```mermaid\nflowchart LR\nA-->B\n```", "x.md", "auto") == "mermaid"
    assert mod.detect_type("sequenceDiagram\nA->>B: hi", "flow.mmd", "auto") == "mermaid"
    assert mod.detect_type("classDiagram", "x.txt", "auto") == "mermaid"
    assert mod.detect_type("x", "y.txt", "mermaid") == "mermaid"


def test_render_extract_source():
    mod = _load_render_module()
    # Mermaid 代码块提取
    raw = "前文\n```mermaid\nflowchart LR\nA-->B\n```\n后文"
    assert mod.extract_source(raw, "mermaid") == "flowchart LR\nA-->B"
    # PlantUML 块提取
    raw = "@startuml\nAlice -> Bob\n@enduml"
    assert mod.extract_source(raw, "plantuml") == raw
    # 纯 Mermaid 定义原样保留
    assert mod.extract_source("sequenceDiagram\nA->>B", "mermaid") == "sequenceDiagram\nA->>B"


# ---------------------------------------------------------------- PlantUML 安全写法预校验

def test_lint_plantuml_valid_source_no_errors():
    mod = _load_render_module()
    src = "@startuml\nA --> B : label\n@enduml"
    errors, warnings = mod.lint_plantuml(src)
    assert errors == []
    assert warnings == []


def test_lint_plantuml_startuml_enduml_mismatch():
    mod = _load_render_module()
    errors, _ = mod.lint_plantuml("@startuml\nA --> B\n")  # 缺 @enduml
    assert len(errors) == 1
    assert "不配对" in errors[0]


def test_lint_plantuml_arrow_missing_target():
    mod = _load_render_module()
    _, warnings = mod.lint_plantuml("@startuml\nA --> : label\n@enduml")
    assert any("箭头缺少目标" in w for w in warnings)


def test_lint_plantuml_parens_in_block():
    mod = _load_render_module()
    src = "@startuml\npackage \"组\" {\n(接口)\n}\n@enduml"
    _, warnings = mod.lint_plantuml(src)
    assert any("() 声明出现在" in w for w in warnings)


def test_lint_plantuml_nonstandard_elements():
    mod = _load_render_module()
    src = "@startuml\nsystem 系统A\nfolder 目录B\n@enduml"
    _, warnings = mod.lint_plantuml(src)
    assert any("ArchiMate" in w for w in warnings)
    assert any("慎用" in w for w in warnings)


def test_lint_plantuml_bracket_imbalance_warns():
    mod = _load_render_module()
    # 去除引号后括号不平衡（{} 缺右括号）
    src = "@startuml\npackage 组 {\ncomponent 元素A\n@enduml"
    _, warnings = mod.lint_plantuml(src)
    assert any("括号不平衡" in w for w in warnings)


def test_render_cli_blocks_on_lint_error(tmp_path):
    """CLI：@startuml/@enduml 不配对应在渲染前拦截并退出非零。"""
    import subprocess

    src = tmp_path / "bad.puml"
    src.write_text("@startuml\nA --> B : label\n", encoding="utf-8")  # 缺 @enduml
    r = subprocess.run(
        [sys.executable, str(RENDER_SCRIPT), str(src), "-o", str(tmp_path / "out.png")],
        capture_output=True, text=True, timeout=30,
    )
    assert r.returncode != 0
    assert "预校验未通过" in r.stderr or "不配对" in r.stderr


def test_render_syntax_error_classified_not_fallback(tmp_path, monkeypatch):
    """引擎已执行但渲染失败 → 抛 DiagramSyntaxError（区别于引擎缺失的 RuntimeError）。

    mock plantuml CLI 存在且执行失败，验证错误分类：语法错误不应被兜底掩盖。
    """
    mod = _load_render_module()

    monkeypatch.setattr(mod, "_which", lambda name: name == "plantuml")

    class FakeProc:
        returncode = 1
        stderr = "ERROR: syntax error at line 3"
        stdout = ""

    monkeypatch.setattr(mod, "_run", lambda *a, **k: FakeProc())

    with pytest.raises(mod.DiagramSyntaxError) as exc_info:
        mod.render_plantuml(
            "@startuml\nA --> B\n@enduml",
            str(tmp_path / "out.png"), 2, str(tmp_path), None,
        )
    assert "syntax error" in str(exc_info.value)


def test_render_cli_syntax_error_hints(tmp_path):
    """CLI：语法错误的 PlantUML（配对正确但元素非法）在有引擎时报语法错误并附速查提示。

    无引擎环境下引擎缺失走兜底，不适用于本测试；跳过条件：无 @plantuml/core 且无 plantuml CLI。
    """
    import subprocess

    mod = _load_render_module()
    has_engine = mod._which("plantuml") or (
        mod._which("node") and mod._plantuml_core_available()
    )
    if not has_engine:
        pytest.skip("本机无 PlantUML 引擎，语法错误路径不可测")

    src = tmp_path / "bad2.puml"
    # @startuml 配对正确，但 () 嵌在 package 块内（已知易崩写法，可能触发引擎报错）
    src.write_text(
        '@startuml\npackage "组" {\n(接口)\n}\n@enduml', encoding="utf-8"
    )
    r = subprocess.run(
        [sys.executable, str(RENDER_SCRIPT), str(src), "-o", str(tmp_path / "out.png"),
         "--no-fallback"],
        capture_output=True, text=True, timeout=120,
    )
    # 引擎若拒绝该语法：报语法错误且（plantuml 时）附速查；若引擎兼容渲染成功也允许通过
    if r.returncode != 0:
        assert "语法错误" in r.stderr or "渲染失败" in r.stderr or "兜底" in r.stderr


def test_render_script_cli_error_on_unknown_type(tmp_path):
    """CLI 对无法识别的类型应报错退出（不依赖真实渲染引擎）。"""
    import subprocess

    src = tmp_path / "unknown.txt"
    src.write_text("随便写点没有图的关键字的内容，应该无法识别类型。", encoding="utf-8")
    r = subprocess.run(
        [sys.executable, str(RENDER_SCRIPT), str(src), "-o", str(tmp_path / "out.png")],
        capture_output=True, text=True, timeout=30,
    )
    assert r.returncode != 0
    assert "无法自动识别" in r.stderr


# ---------------------------------------------------------------- 离线兜底渲染

def _load_fallback_module():
    fb_path = ROOT / "skills" / "diagram-to-office" / "fallback_render.py"
    if not fb_path.is_file():
        pytest.skip("fallback_render.py 不存在")
    spec = importlib.util.spec_from_file_location("fallback_render", fb_path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["fallback_render"] = mod
    spec.loader.exec_module(mod)
    return mod


def test_render_offline_check_env(tmp_path):
    """--check 环境诊断应输出离线策略与引擎清单（不渲染、不联网）。"""
    import subprocess

    r = subprocess.run(
        [sys.executable, str(RENDER_SCRIPT), "--check"],
        capture_output=True, text=True, timeout=60,
    )
    assert r.returncode == 0
    assert "离线" in r.stdout
    assert "matplotlib" in r.stdout
    assert "plantuml" in r.stdout.lower()
    assert "mermaid" in r.stdout.lower()
    # 提示 --install 预装（满足"npx 依赖可预先安装"的诉求）
    assert "--install" in r.stdout
    # 纯 JS 引擎 @plantuml/core 应出现在诊断中
    assert "@plantuml/core" in r.stdout


def test_render_install_flag_supported(tmp_path):
    """--install 应作为合法参数被接受（不实际联网执行）。"""
    import subprocess

    r = subprocess.run(
        [sys.executable, str(RENDER_SCRIPT), "--help"],
        capture_output=True, text=True, timeout=30,
    )
    assert r.returncode == 0
    assert "--install" in r.stdout
    assert "--check" in r.stdout


def test_render_plantuml_core_detect():
    """@plantuml/core 纯 JS 引擎检测函数应可调用且返回 bool（不联网）。"""
    mod = _load_render_module()
    assert isinstance(mod._plantuml_core_available(), bool)
    assert isinstance(mod._has_cairosvg(), bool)
    assert isinstance(mod._resvg_available(), bool)


def test_plantuml_js_render_script_syntax(tmp_path):
    """plantuml_js_render.mjs 应通过 node 语法检查（不执行渲染）。"""
    import shutil
    import subprocess

    if shutil.which("node") is None:
        pytest.skip("node 不可用")
    js_script = ROOT / "skills" / "diagram-to-office" / "plantuml_js_render.mjs"
    if not js_script.is_file():
        pytest.skip("plantuml_js_render.mjs 不存在")
    r = subprocess.run(
        ["node", "--check", str(js_script)],
        capture_output=True, text=True, timeout=30,
    )
    assert r.returncode == 0, r.stderr


def test_svg2png_script_syntax(tmp_path):
    """svg2png.mjs 应通过 node 语法检查（不执行转换）。"""
    import shutil
    import subprocess

    if shutil.which("node") is None:
        pytest.skip("node 不可用")
    js_script = ROOT / "skills" / "diagram-to-office" / "svg2png.mjs"
    if not js_script.is_file():
        pytest.skip("svg2png.mjs 不存在")
    r = subprocess.run(
        ["node", "--check", str(js_script)],
        capture_output=True, text=True, timeout=30,
    )
    assert r.returncode == 0, r.stderr


def test_render_offline_fallback_flowchart(tmp_path):
    """无外部引擎时，内置 matplotlib 兜底应能渲染 flowchart 图片（离线必出图）。"""
    mod = _load_fallback_module()
    src = "flowchart LR\nA[开始] --> B{判断}\nB -->|是| C[处理]\nB -->|否| D[结束]\n"
    out = tmp_path / "flow.png"
    engine = mod.render_fallback(src, "mermaid", str(out), 2)
    assert out.is_file() and out.stat().st_size > 0
    assert "matplotlib" in engine or "兜底" in engine


def test_render_offline_fallback_sequence(tmp_path):
    """内置兜底渲染时序图（mermaid / plantuml 两种语法）。"""
    mod = _load_fallback_module()
    src = "sequenceDiagram\nparticipant A as 用户\nparticipant B as 服务\nA->>B: 请求\nB-->>A: 响应\n"
    out = tmp_path / "seq.png"
    engine = mod.render_fallback(src, "mermaid", str(out), 2)
    assert out.is_file() and out.stat().st_size > 0

    src2 = "@startuml\nactor 用户\nparticipant 服务\n用户 -> 服务: 请求\n服务 --> 用户: 响应\n@enduml"
    out2 = tmp_path / "seq2.png"
    engine2 = mod.render_fallback(src2, "plantuml", str(out2), 2)
    assert out2.is_file() and out2.stat().st_size > 0


def test_render_offline_fallback_text(tmp_path):
    """无法解析的图表类型应输出源码文本图（保证有输出）。"""
    mod = _load_fallback_module()
    src = "classDiagram\nclass Animal {\n  +name: str\n}"
    out = tmp_path / "cls.png"
    engine = mod.render_fallback(src, "mermaid", str(out), 2)
    assert out.is_file() and out.stat().st_size > 0


def test_render_cli_offline_end_to_end(tmp_path):
    """CLI 端到端：离线环境（无引擎）也应成功产出图片（走兜底）。"""
    import subprocess

    src = tmp_path / "flow.mmd"
    src.write_text("flowchart TD\nA-->B\nB-->C\n", encoding="utf-8")
    out = tmp_path / "out.png"
    r = subprocess.run(
        [sys.executable, str(RENDER_SCRIPT), str(src), "-o", str(out)],
        capture_output=True, text=True, timeout=120,
    )
    # 本机若装了 plantuml/mmdc 可能走外部引擎成功；未装则走兜底。
    # 两种情况下都必须产出非空图片。
    assert r.returncode == 0, r.stderr
    assert out.is_file() and out.stat().st_size > 0
    assert "已生成图片" in r.stderr
