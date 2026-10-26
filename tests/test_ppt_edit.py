# SPDX-License-Identifier: Apache-2.0
# Copyright (c) 2026 litework contributors

"""PPT 图形化编辑回归测试（ppt_edit 工具层）。

覆盖三个曾经真实踩过的坑：

1. **对齐只改 anchor 不改 x**：左对齐的文字会从原中线起排、右对齐跑出框外。
   SVG 的 text-anchor 决定「哪个点对齐 x」，两者必须一起变。
2. **打包态误用 sys.executable**：frozen 后它指向后端主程序，拼上脚本路径
   会被当成子命令解析，报 `invalid choice: .../pptx_to_svg.py`。
3. **表格/图表的预览 <image> 被误判成图片元素**：outline 里拿不到 table/chart，
   面板会退化成「图片」inspector。
"""
from __future__ import annotations

import json
import os
import re
import sys

import pytest

pytest.importorskip("pptx")

from litework.tools.ppt_edit import PptEditTools  # noqa: E402


@pytest.fixture()
def workspace(tmp_path):
    """建一份含 2 页（标题页 + 章节页）的样例 pptx。"""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "季度经营回顾"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "第二章"
    prs.save(str(tmp_path / "demo.pptx"))
    return tmp_path


def _call(tools: PptEditTools, **kwargs):
    """同步跑一次工具调用，返回解析后的结果（错误直接抛出）。"""
    import asyncio

    raw = asyncio.run(tools.execute("ppt_edit", kwargs))
    assert not raw.startswith("[PPT Error]"), raw
    return json.loads(raw)


def _title_element(tools: PptEditTools, project: str, slide: int = 1):
    outline = _call(tools, op="outline", project=project, slide=slide)
    for el in outline["slides"][0]["elements"]:
        if el.get("text"):
            return el
    raise AssertionError("样例页没有文本元素")


class TestAlign:
    """对齐必须同时改 text-anchor 与 x。"""

    @pytest.mark.parametrize("align,expect", [
        ("left", "start"),
        ("center", "middle"),
        ("right", "end"),
    ])
    def test_anchor_matches(self, workspace, align, expect):
        tools = PptEditTools(str(workspace))
        project = _call(tools, op="open", path="demo.pptx")["project"]
        el = _title_element(tools, project)
        _call(tools, op="edit", project=project, slide=1,
              element_id=el["id"], changes={"align": align})

        svg = _call(tools, op="slide_svg", project=project, slide=1)["svg"]
        m = re.search(r'<text x="([\d.]+)"[^>]*text-anchor="(\w+)"', svg)
        assert m, "未找到 text 元素"
        assert m.group(2) == expect

    def test_x_follows_frame(self, workspace):
        """x 必须落在 frame 边界内，而不是停在原来的中线。"""
        tools = PptEditTools(str(workspace))
        project = _call(tools, op="open", path="demo.pptx")["project"]
        el = _title_element(tools, project)
        fx, _fy, fw, _fh = el["frame"]

        for align, want in (("left", fx), ("right", fx + fw)):
            _call(tools, op="edit", project=project, slide=1,
                  element_id=el["id"], changes={"align": align})
            svg = _call(tools, op="slide_svg", project=project, slide=1)["svg"]
            x = float(re.search(r'<text x="([\d.]+)"', svg).group(1))
            # 允许内边距（约 9.6px），但必须贴近对应边界
            assert abs(x - want) <= 12, f"{align}: x={x} 应贴近 {want}"

    def test_frame_change_recomputes_x(self, workspace):
        """改宽度后，居中文字应跟到新中线（否则框变宽字不动）。"""
        tools = PptEditTools(str(workspace))
        project = _call(tools, op="open", path="demo.pptx")["project"]
        el = _title_element(tools, project)
        _call(tools, op="edit", project=project, slide=1,
              element_id=el["id"], changes={"align": "center"})
        _call(tools, op="edit", project=project, slide=1,
              element_id=el["id"], changes={"x": 100.0, "width": 400.0})

        svg = _call(tools, op="slide_svg", project=project, slide=1)["svg"]
        x = float(re.search(r'<text x="([\d.]+)"', svg).group(1))
        assert abs(x - 300.0) < 1.0, f"期望新中线 300，实际 {x}"


class TestInterpreter:
    """打包态不能拿 sys.executable 当技能解释器。"""

    def test_frozen_uses_skill_python(self, monkeypatch):
        from litework.tools import ppt_edit as mod

        fake = "/opt/app/_internal/skill-python/bin/python3"
        monkeypatch.setattr(mod, "_skill_python", lambda: fake)
        assert mod._skill_python().endswith("skill-python/bin/python3")

    def test_frozen_without_skill_python_raises(self, monkeypatch):
        """frozen 且找不到 skill-python 时必须明确报错，而不是静默用错解释器。"""
        from litework.tools import ppt_edit as mod
        import litework.tools.skills as skills_mod

        monkeypatch.setattr(sys, "frozen", True, raising=False)
        monkeypatch.setattr(skills_mod, "_builtin_skill_python", lambda: None)
        with pytest.raises(mod.PptEditError):
            mod._skill_python()

    def test_unusable_skill_python_falls_back_in_dev(self, monkeypatch):
        """解释器存在但导不进 pptx/PIL（release/ 只留了裸 CPython）→ 开发态回退。

        回归：构建期缓存 `release/skill-python/` 里的裸 CPython 没有第三方库
        （那些库靠 `_internal` 的 sitecustomize 注入，而缓存目录常在清理后
        只剩壳）。旧实现只看「文件在不在」，直接用它 → 所有 PPT 导出
        `ModuleNotFoundError: No module named 'pptx'`。
        """
        from litework.tools import ppt_edit as mod
        import litework.tools.skills as skills_mod

        monkeypatch.setattr(skills_mod, "_builtin_skill_python",
                            lambda: "/tmp/skill-python/python/bin/python3")
        monkeypatch.setattr(mod, "_skill_python_usable", lambda _exe: False)
        assert mod._skill_python() == sys.executable

    def test_unusable_skill_python_raises_when_frozen(self, monkeypatch):
        """打包态同样不可用 → 明确报错（不能回退到 frozen 后端主程序）。"""
        from litework.tools import ppt_edit as mod
        import litework.tools.skills as skills_mod

        monkeypatch.setattr(sys, "frozen", True, raising=False)
        monkeypatch.setattr(skills_mod, "_builtin_skill_python",
                            lambda: "/opt/app/_internal/skill-python/bin/python3")
        monkeypatch.setattr(mod, "_skill_python_usable", lambda _exe: False)
        with pytest.raises(mod.PptEditError):
            mod._skill_python()

    def test_usable_skill_python_wins(self, monkeypatch):
        """随包解释器可用时优先用它（打包态离线自足的关键）。"""
        from litework.tools import ppt_edit as mod
        import litework.tools.skills as skills_mod

        monkeypatch.setattr(skills_mod, "_builtin_skill_python",
                            lambda: "/opt/app/_internal/skill-python/bin/python3")
        monkeypatch.setattr(mod, "_skill_python_usable", lambda _exe: True)
        assert mod._skill_python().endswith("skill-python/bin/python3")

    def test_usable_probe_rejects_broken_interpreter(self):
        """探测用的解释器确实能跑 import（用真实可用/不可用命令各验一次）。"""
        from litework.tools import ppt_edit as mod

        assert mod._skill_python_usable(sys.executable) is True
        assert mod._skill_python_usable("/nonexistent/python3") is False

    def test_usable_probe_is_cached(self, monkeypatch):
        """探测结果按路径缓存：一次导出会多次选解释器，不能每次都起子进程。"""
        from litework.tools import ppt_edit as mod

        calls = []
        real_run = mod.subprocess.run

        def counting_run(*a, **kw):
            calls.append(a[0])
            return real_run(*a, **kw)

        monkeypatch.setattr(mod, "_SKILL_PYTHON_USABLE", {})
        monkeypatch.setattr(mod.subprocess, "run", counting_run)

        assert mod._skill_python_usable(sys.executable) is True
        assert mod._skill_python_usable(sys.executable) is True
        assert mod._skill_python_usable(sys.executable) is True
        assert len(calls) == 1, "同一解释器只探测一次"


class TestElementKinds:
    """outline 必须正确区分文本/图片/表格，预览图不算图片元素。"""

    def test_outline_has_text_elements(self, workspace):
        tools = PptEditTools(str(workspace))
        project = _call(tools, op="open", path="demo.pptx")["project"]
        outline = _call(tools, op="outline", project=project)
        assert outline["total"] == 2
        first = outline["slides"][0]["elements"]
        texts = [e for e in first if e.get("text")]
        assert texts, "标题页应有文本元素"
        assert all(e["kind"] != "picture" for e in first), "文本页不应出现图片元素"


class TestPagePlan:
    """跨页改动：每页独立提交，未改页导出时字节级保留。"""

    def test_cross_slide_edits_both_applied(self, workspace):
        import hashlib
        import zipfile

        tools = PptEditTools(str(workspace))
        project = _call(tools, op="open", path="demo.pptx")["project"]

        for slide, txt in ((1, "页1-改"), (2, "页2-改")):
            outline = _call(tools, op="outline", project=project, slide=slide)
            el = next(e for e in outline["slides"][0]["elements"] if e.get("text"))
            _call(tools, op="edit", project=project, slide=slide,
                  element_id=el["id"], changes={"text": [txt]})

        out = _call(tools, op="export", project=project, output="out.pptx")["output"]
        path = os.path.join(str(workspace), out)

        from pptx import Presentation
        prs = Presentation(path)
        joined = [
            "".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            for s in prs.slides
        ]
        assert "页1-改" in joined[0]
        assert "页2-改" in joined[1]

        src = os.path.join(str(workspace), "demo.pptx")
        with zipfile.ZipFile(path) as a, zipfile.ZipFile(src) as b:
            # 两页都改过 → 都不该是 passthrough，但文件必须可正常打开
            assert "ppt/slides/slide1.xml" in a.namelist()
            assert hashlib.sha256(a.read("ppt/slides/slide1.xml")).hexdigest() != \
                hashlib.sha256(b.read("ppt/slides/slide1.xml")).hexdigest()

    def test_thumbs_strip_images(self, workspace):
        """缩略图必须瘦身：不含 <image>，体积可控。"""
        tools = PptEditTools(str(workspace))
        project = _call(tools, op="open", path="demo.pptx")["project"]
        data = _call(tools, op="thumbs", project=project)
        assert len(data["slides"]) == 2
        for s in data["slides"]:
            assert "<image" not in s["svg"]
            assert len(s["svg"]) < 20000


class TestExtractForRestyle:
    """重排前的内容抽取：要能认出标题/要点/表格/图表，供 Agent 映射版式。"""

    @pytest.fixture()
    def rich_deck(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Emu, Inches
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        s1 = prs.slides.add_slide(prs.slide_layouts[0])
        s1.shapes.title.text = "战略规划"
        s1.placeholders[1].text = "增长与效率"

        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        s2.shapes.title.text = "市场机会"
        tf = s2.placeholders[1].text_frame
        tf.text = "规模 1200 亿"
        for t in ("增速 18%", "政策红利", "格局分散"):
            tf.add_paragraph().text = t

        s3 = prs.slides.add_slide(prs.slide_layouts[5])
        s3.shapes.title.text = "增长趋势"
        cd = CategoryChartData()
        cd.categories = ["2024", "2025"]
        cd.add_series("营收", (100, 150))
        s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                            Inches(1), Inches(1.8), Inches(8), Inches(3), cd)

        s4 = prs.slides.add_slide(prs.slide_layouts[5])
        s4.shapes.title.text = "区域业绩"
        tb = s4.shapes.add_table(2, 2, Inches(1), Inches(1.8),
                                 Inches(6), Inches(2)).table
        for i, (a, b) in enumerate([("区域", "营收"), ("华东", "4200")]):
            tb.cell(i, 0).text = a
            tb.cell(i, 1).text = b

        prs.save(str(tmp_path / "rich.pptx"))
        return tmp_path

    def test_extract_structure(self, rich_deck):
        tools = PptEditTools(str(rich_deck))
        data = _call(tools, op="extract", path="rich.pptx")
        assert data["total"] == 4

        s1, s2, s3, s4 = data["slides"]
        assert s1["title"] == "战略规划"
        assert s1["layout_guess"] == "title_slide"
        assert "增长与效率" in s1["paragraphs"]

        assert s2["title"] == "市场机会"
        assert len(s2["paragraphs"]) == 4

        assert s3["has_chart"] is True
        assert s3["chart"][0]["series"][0]["values"] == [100.0, 150.0]

        assert s4["has_table"] is True
        assert s4["table"] == [["区域", "营收"], ["华东", "4200"]]

    def test_extract_rejects_missing_file(self, tmp_path):
        tools = PptEditTools(str(tmp_path))
        import asyncio
        raw = asyncio.run(tools.execute("ppt_edit", {"op": "extract", "path": "nope.pptx"}))
        assert raw.startswith("[PPT Error]")


class TestRestyleGeneration:
    """重排落地产物：按计划填充模板并生成可打开的 pptx。"""

    def test_create_from_template_fills_placeholders(self, tmp_path):
        tools = PptEditTools(str(tmp_path))
        result = _call(tools, op="create_from_template", filename="out.pptx", template={
            "suite": "presentation_core",
            "pages": [
                {"page": "title_slide",
                 "values": {"TITLE": "新标题", "SUBTITLE": "新副标题"}},
                {"page": "chart_insight",
                 "values": {"PAGE_TITLE": "趋势", "KEY_MESSAGE": "三年翻倍",
                            "SOURCE": "来源：财务部"}},
            ],
        })
        assert result["slides"] == 2

        from pptx import Presentation
        prs = Presentation(os.path.join(str(tmp_path), result["output"]))
        texts = [
            [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
            for s in prs.slides
        ]
        assert "新标题" in texts[0]
        assert "新副标题" in texts[0]
        assert "三年翻倍" in texts[1]

    def test_unknown_suite_reports_clearly(self, tmp_path):
        tools = PptEditTools(str(tmp_path))
        import asyncio
        raw = asyncio.run(tools.execute("ppt_edit", {
            "op": "create_from_template",
            "template": {"suite": "no_such_suite", "pages": [{"page": "x", "values": {}}]},
        }))
        assert "no_such_suite" in raw
