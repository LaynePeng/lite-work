# SPDX-License-Identifier: Apache-2.0
# Copyright (c) 2026 lite-work contributors

"""`/api/ppt/plan`：提示词 → 逐页方案（一次模型调用，不落盘）。

关键契约：

1. 返回的 `pages` 可直接喂给 `/api/ppt/create`（page 是模板 slug，占位符已校验）；
2. 模型输出不合规时**不 500**，降级成按版式排的骨架页并在 `error` 里说明原因；
3. 显式 `fallback=false` 时才把失败当错误返回（给需要严格语义的调用方）。
"""
from __future__ import annotations

import json

import pytest
from fastapi.testclient import TestClient

from litework.app import AgentApp
from litework.server.app import create_app
from tests.conftest import MockLLMAdapter

pytest.importorskip("pptx")


@pytest.fixture()
def client_and_app(tmp_path):
    app = AgentApp(workspace=str(tmp_path), config_dir=str(tmp_path / ".lite-work"))
    fast = create_app(app, token=None)
    with TestClient(fast) as client:
        yield client, app


GOOD_PLAN = json.dumps({"pages": [
    {"page": "title_slide", "values": {"TITLE": "Q3 复盘", "SUBTITLE": "管理层汇报"}},
    {"page": "title_content", "values": {"PAGE_TITLE": "营收", "CONTENT_AREA": "增长 12%\n毛利提升"}},
]}, ensure_ascii=False)


def test_plan_returns_pages_usable_by_create(client_and_app):
    client, app = client_and_app
    app._mock_adapter = MockLLMAdapter([(GOOD_PLAN, [])])

    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core", "brief": "给管理层做 Q3 复盘",
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["fallback"] is False
    assert [p["page"] for p in body["pages"]] == ["title_slide", "title_content"]
    assert body["pages"][1]["values"]["CONTENT_AREA"] == "增长 12%\n毛利提升"

    # 规划结果原样喂给 create 必须能生成
    c = client.post("/api/ppt/create", json={
        "template": {"suite": "presentation_core",
                     "pages": [{"page": p["page"], "values": p["values"]} for p in body["pages"]]},
        "filename": "planned.pptx",
    })
    assert c.status_code == 200, c.text
    assert c.json()["slides"] == 2


def test_plan_falls_back_when_model_output_is_junk(client_and_app):
    client, app = client_and_app
    app._mock_adapter = MockLLMAdapter([("我觉得应该先做个封面", [])])

    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core", "brief": "做一份复盘", "page_hint": 5,
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["fallback"] is True
    assert body["error"], "降级必须说明原因"
    assert len(body["pages"]) == 5, "兜底应遵守 page_hint"
    assert all(p["page"] for p in body["pages"])


def test_plan_strict_mode_returns_422_on_junk(client_and_app):
    client, app = client_and_app
    app._mock_adapter = MockLLMAdapter([("不是 JSON", [])])

    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core", "brief": "做一份复盘", "fallback": False,
    })
    assert r.status_code == 422
    assert r.json()["detail"]


def test_plan_rejects_unknown_suite(client_and_app):
    client, app = client_and_app
    client.post("/api/workspace", json={"path": str(app.workspace)})
    r = client.post("/api/ppt/plan", json={"suite": "nope", "brief": "x"})
    assert r.status_code == 422
    assert "nope" in r.json()["detail"]


def test_plan_requires_brief_without_source(client_and_app):
    client, app = client_and_app
    app._mock_adapter = MockLLMAdapter([(GOOD_PLAN, [])])
    r = client.post("/api/ppt/plan", json={"suite": "presentation_core", "brief": "  "})
    assert r.status_code == 422


def test_plan_restyle_uses_source_slides(client_and_app, tmp_path):
    """重排：给了 source_path 就该把旧稿内容写进提示词。"""
    client, app = client_and_app

    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "旧标题"
    s1.placeholders[1].text = "旧要点"
    prs.save(str(tmp_path / "old.pptx"))

    seen: list[str] = []

    class Recorder(MockLLMAdapter):
        async def chat_stream(self, messages, tools, events=None):
            seen.extend(m.content or "" for m in messages)
            return await super().chat_stream(messages, tools, events)

    app._mock_adapter = Recorder([(GOOD_PLAN, [])])
    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core", "source_path": "old.pptx",
    })
    assert r.status_code == 200, r.text
    joined = "\n".join(seen)
    assert "旧标题" in joined and "旧要点" in joined


# ---------------------------------------------------------------- 技能不可用

def _break_skill(monkeypatch):
    from litework.tools import ppt_edit as pe

    def boom():
        raise pe.PptEditError("找不到内置 ppt-master 技能（skills/ppt-master）。")

    monkeypatch.setattr(pe, "_skill_dir", boom)


def test_templates_reports_unavailable_instead_of_empty(client_and_app, monkeypatch):
    """技能缺失时 /api/ppt/templates 必须说明原因，而不是返回「空模板库」。"""
    client, app = client_and_app
    _break_skill(monkeypatch)

    r = client.get("/api/ppt/templates")
    assert r.status_code == 200
    body = r.json()
    assert body["layouts"] == []
    assert body["available"]["ready"] is False
    assert "ppt-master" in body["available"]["reason"]
    assert body["available"]["hint"]


def test_plan_returns_409_when_skill_missing(client_and_app, monkeypatch):
    """技能缺失是「能力不可用」，不是「参数错误」——必须 409 且给修复建议。"""
    client, app = client_and_app
    _break_skill(monkeypatch)

    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core", "brief": "做一份复盘"})
    assert r.status_code == 409, r.text
    assert "安装" in r.json()["detail"] or "ppt-master" in r.json()["detail"]


def test_create_returns_409_when_skill_missing(client_and_app, monkeypatch):
    client, app = client_and_app
    _break_skill(monkeypatch)

    r = client.post("/api/ppt/create", json={"template": {
        "suite": "presentation_core", "pages": [{"page": "title_slide", "values": {}}]}})
    assert r.status_code == 409, r.text


def test_open_returns_409_when_skill_missing(client_and_app, monkeypatch, tmp_path):
    """编辑器入口同理：不能让用户对着原始 [PPT Error] 文本猜。"""
    client, app = client_and_app
    _break_skill(monkeypatch)
    (tmp_path / "x.pptx").write_bytes(b"not a real pptx")

    r = client.post("/api/ppt/open", json={"path": "x.pptx"})
    assert r.status_code == 409, r.text


# ---------------------------------------------------------------- AI 配图接线

def test_images_endpoint_reports_backends_from_chat_providers(client_and_app):
    """生图后端凭据来自已配置的聊天供应商，不该要求用户再填一遍。"""
    client, app = client_and_app
    app.llm_registry.providers["openai"]["api_key"] = "sk-test-openai"

    r = client.get("/api/ppt/images")
    assert r.status_code == 200
    body = r.json()
    hit = next((b for b in body["backends"] if b["backend"] == "openai"), None)
    assert hit is not None, body
    assert hit["from_llm_provider"] is True


def test_images_endpoint_empty_without_credentials(client_and_app, monkeypatch):
    client, app = client_and_app
    for pid in ("openai", "qwen", "glm"):
        app.llm_registry.providers.get(pid, {})["api_key"] = ""
    for key in ("OPENAI_API_KEY", "QWEN_API_KEY", "DASHSCOPE_API_KEY",
                "ZHIPU_API_KEY", "BIGMODEL_API_KEY", "GEMINI_API_KEY"):
        monkeypatch.delenv(key, raising=False)
    body = client.get("/api/ppt/images").json()
    assert body["backends"] == []


def test_create_runs_image_generation_and_fills_slots(client_and_app, monkeypatch, tmp_path):
    """配图成功后：图片值进了 values，最终 pptx 里有真实图片。"""
    import zlib
    import struct

    client, app = client_and_app
    app.llm_registry.providers["openai"]["api_key"] = "sk-test-openai"

    def png(w, h, rgb):
        raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

        def chunk(t, d):
            c = t + d
            return struct.pack(">I", len(d)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)

        return (b"\x89PNG\r\n\x1a\n"
                + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))

    calls = []

    def fake_generate(**kwargs):
        calls.append(kwargs)
        path = os.path.join(str(tmp_path), "素材", "ppt-images", "p01-image1.png")
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "wb") as fh:
            fh.write(png(1024, 576, (90, 30, 160)))
        return {"ok": True, "path": path, "backend": "openai", "error": "", "log": ""}

    from litework.tools import ppt_images
    monkeypatch.setattr(ppt_images, "generate", fake_generate)
    # 图库通道也 pin 住：生图成功时不该走到它，但避免任何真实网络
    monkeypatch.setattr(ppt_images, "fetch_stock_tiered", lambda **kw: {
        "ok": False, "path": "", "provider": "", "attribution": "", "error": "stub"})

    r = client.post("/api/ppt/create", json={
        "template": {"suite": "presentation_core", "pages": [{
            "page": "title_picture", "values": {"PAGE_TITLE": "封面"}}]},
        "filename": "withimg.pptx",
        "images": {"1:IMAGE_1": "modern office, corporate photography"},
    })
    assert r.status_code == 200, r.text
    assert calls, "没有调用生图"
    # 画面描述被加上「不许出现文字」的约束
    assert "no text" in calls[0]["prompt"].lower()

    from pptx import Presentation
    prs = Presentation(str(tmp_path / "withimg.pptx"))
    pics = [sh for s in prs.slides for sh in s.shapes
            if sh.shape_type == 13 or (sh.shape_type == 6 and any(
                c.shape_type == 13 for c in sh.shapes))]
    assert pics, "pptx 里没有图片"


def test_create_image_failure_only_warns(client_and_app, monkeypatch):
    """单张配图失败不能把整份稿子废掉。"""
    client, app = client_and_app
    app.llm_registry.providers["openai"]["api_key"] = "sk-test-openai"

    from litework.tools import ppt_images
    monkeypatch.setattr(ppt_images, "generate", lambda **kw: {
        "ok": False, "path": "", "backend": "openai",
        "error": "额度不足", "log": ""})
    # 生图失败会逐张降级到图库；这里让图库也没结果，验证「都失败仍出稿」
    monkeypatch.setattr(ppt_images, "fetch_stock_tiered", lambda **kw: {
        "ok": False, "path": "", "provider": "", "attribution": "",
        "error": "图库也没找到"})

    r = client.post("/api/ppt/create", json={
        "template": {"suite": "presentation_core", "pages": [{
            "page": "title_picture", "values": {"PAGE_TITLE": "封面"}}]},
        "filename": "noimg.pptx",
        "images": {"1:IMAGE_1": "a modern office"},
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["slides"] == 1
    assert any("额度不足" in w for w in body["warnings"]), body["warnings"]


# ---------------------------------------------------------------- B 组能力

def test_preview_renders_filled_svg(client_and_app):
    """B1：规划预览——填好的文案要出现在 SVG 里。"""
    client, app = client_and_app
    r = client.post("/api/ppt/preview", json={
        "suite": "presentation_core", "page": "three_card",
        "values": {"PAGE_TITLE": "三个增长点",
                   "CARD_1": "续费 92%", "CARD_2": "海外翻倍", "CARD_3": "人效 +18%"}})
    assert r.status_code == 200
    body = r.json()
    assert body["ok"]
    assert "续费 92%" in body["svg"]
    assert body["svg"].startswith("<?xml") or body["svg"].startswith("<svg")


def test_preview_rejects_unknown_page(client_and_app):
    client, app = client_and_app
    r = client.post("/api/ppt/preview", json={
        "suite": "presentation_core", "page": "nope"})
    assert r.status_code == 422


def test_rewrite_page_via_plan(client_and_app):
    """B3：单页重写——只返回这一页的新文案。"""
    client, app = client_and_app
    app._mock_adapter = MockLLMAdapter([(json.dumps({
        "values": {"PAGE_TITLE": "更犀利的标题", "CONTENT_AREA": "新要点一\n新要点二"},
    }, ensure_ascii=False), [])])

    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core",
        "rewrite": {"page": "title_content",
                    "values": {"PAGE_TITLE": "旧标题", "CONTENT_AREA": "旧内容"},
                    "instruction": "标题更犀利，正文改成两条新要点"}})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rewrite"] is True
    assert body["values"]["PAGE_TITLE"] == "更犀利的标题"
    assert body["values"]["CONTENT_AREA"] == "新要点一\n新要点二"


def test_rewrite_filters_unknown_keys(client_and_app):
    """模型编造的占位符键必须被过滤。"""
    client, app = client_and_app
    app._mock_adapter = MockLLMAdapter([(json.dumps({
        "values": {"PAGE_TITLE": "ok", "BOGUS_KEY": "x"},
    }, ensure_ascii=False), [])])
    r = client.post("/api/ppt/plan", json={
        "suite": "presentation_core",
        "rewrite": {"page": "title_content", "values": {},
                    "instruction": "改一下"}})
    body = r.json()
    assert "BOGUS_KEY" not in body["values"]
    assert body["values"]["PAGE_TITLE"] == "ok"


def test_export_pdf_reports_soffice_missing(client_and_app, tmp_path, monkeypatch):
    """B2：无 LibreOffice 时 PDF 失败要可读，且 pptx 不作废。"""
    client, app = client_and_app
    from litework.tools import office

    monkeypatch.setattr(office, "_find_soffice", lambda: None)
    r = client.post("/api/ppt/create", json={
        "template": {"suite": "presentation_core", "pages": [{
            "page": "title_slide", "values": {"TITLE": "t"}}]},
        "filename": "p.pptx"})
    assert r.status_code == 200
    project = r.json()["project"]

    e = client.post("/api/ppt/export", json={
        "project": project, "output": "p_out.pptx", "format": "pdf"})
    assert e.status_code == 200
    body = e.json()
    assert body["ok"] is True, "pptx 导出不能因 pdf 失败作废"
    assert "LibreOffice" in (body.get("pdf_error") or ""), body


def test_notes_roundtrip(client_and_app, tmp_path):
    """B4：备注写读 + 清除。"""
    client, app = client_and_app
    r = client.post("/api/ppt/create", json={
        "template": {"suite": "presentation_core", "pages": [{
            "page": "title_slide", "values": {"TITLE": "n"}}]},
        "filename": "notes.pptx"})
    project = r.json()["project"]

    w = client.post("/api/ppt/set_note", json={
        "project": project, "slide": 1, "note": "开场白：三件事"})
    assert w.status_code == 200 and w.json()["ok"]

    g = client.get(f"/api/ppt/notes?project={project}")
    assert g.status_code == 200
    notes = {n["slide"]: n["note"] for n in g.json()["notes"]}
    assert notes[1] == "开场白：三件事"

    # 清除
    c = client.post("/api/ppt/set_note", json={
        "project": project, "slide": 1, "note": ""})
    assert c.json()["cleared"] is True
    g2 = client.get(f"/api/ppt/notes?project={project}")
    assert {n["slide"]: n["note"] for n in g2.json()["notes"]}[1] == ""
